"""
Relatorio Excel formatado, com as cinco abas da secao 6.2.

O painel HTML e para olhar; esta planilha e para trabalhar em cima. Por isso
duas escolhas atravessam o modulo inteiro:

1. Numero sai como NUMERO, nunca como texto formatado. A celula leva o valor
   cru e o formato de exibicao vai no number_format. Gravar '3,18%' como texto
   deixa a planilha bonita e quebra SOMA, MEDIA e ordenacao - que e exatamente
   o que quem abre o arquivo vai querer fazer com ele.

2. Indicador sem dado fica com a celula VAZIA, nao com o texto 'sem dado'. Uma
   palavra no meio de uma coluna numerica contamina a coluna toda; a celula
   vazia e ignorada por MEDIA e some de um grafico, que e o tratamento correto
   para o que nao foi medido. A aba Resumo explica a convencao.

A aba de qualidade nao sai do banco: o banco guarda so a contagem da execucao,
nao a ocorrencia linha a linha. Ela vem da importacao, que e idempotente - rodar
de novo sobre a mesma pasta reproduz as mesmas ocorrencias sem duplicar ordem
nenhuma.
"""

import argparse
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.banco import ErroBanco, conectar, importar
from src.config import (
    CODIGO_SAIDA_ERRO_CONFIG,
    CONFIG_PADRAO,
    Configuracao,
    ErroConfiguracao,
    carregar,
)
from src.indicadores import (
    ErroIndicadores,
    Indicadores,
    ResultadoIndicadores,
    calcular,
    calcular_corte,
    maquinas_na_meta,
    meta_de,
    na_meta,
    produtividade_na_meta,
)
from src.leitura import ErroLeitura
from src.painel import ordens_atrasadas
from src.validacao import ErroValidacao, Ocorrencia

ARQUIVO_PADRAO = "relatorio_producao.xlsx"

# As cinco abas da secao 6.2, na ordem em que ela as declara. Nomeadas aqui em
# vez de literais espalhados por montar_livro: o Excel recusa nome com mais de
# 31 caracteres ou com [ ] : * ? / \ e abre o arquivo com aviso de reparo, e um
# lugar so e um lugar so para conferir isso.
ABAS_ESPERADAS = (
    "Resumo",
    "Por máquina",
    "Por cliente",
    "Atrasos",
    "Qualidade dos dados",
)

# Excel recusa mais que isso numa celula e a planilha abre com aviso de reparo.
MAXIMO_POR_CELULA = 32_000

# Formatos de exibicao. O separador de milhar e decimal sai do idioma do Excel de
# quem abre, e nao daqui: '#.##0' escrito assim e interpretado no locale do
# usuario, que e o que se quer num arquivo que vai circular.
FORMATO_INTEIRO = "#,##0"
FORMATO_UM_DECIMAL = "#,##0.0"
FORMATO_PERCENTUAL = "#,##0.00"
FORMATO_DATA = "DD/MM/YYYY"

# Casas decimais so quando existem. Serve a coluna Valor do Resumo, que mistura
# contagem (678) com indicador (3,18) - um formato fixo mostraria "678,0".
#
# E a unica mascara que NAO pode ir direto para a celula: veja _mascara.
FORMATO_GERAL = "#,##0.##"

CABECALHO_FUNDO = "1F3A44"
CABECALHO_TEXTO = "FFFFFF"


class ErroRelatorio(Exception):
    """Relatorio que nao pode ser montado ou gravado."""


def _texto_seguro(valor) -> str:
    """
    Limpa o texto que vai para uma celula.

    O Excel recusa caractere de controle, e o openpyxl levanta excecao ao
    encontrar um. O 'valor_original' de uma ocorrencia carrega o conteudo cru da
    planilha de origem, entao ele pode trazer qualquer coisa - e o criterio de
    aceite deste bloco e o arquivo abrir sem aviso de reparo.
    """
    if valor is None:
        return ""
    texto = ILLEGAL_CHARACTERS_RE.sub("", str(valor))
    if len(texto) > MAXIMO_POR_CELULA:
        return texto[: MAXIMO_POR_CELULA - 3] + "..."
    return texto


def _mascara(formato: str | None, valor) -> str | None:
    """
    Escolhe a mascara efetiva de uma celula.

    So FORMATO_GERAL depende do valor, e a razao e uma armadilha do codigo de
    formato do Excel: nele o '.' e sempre o separador decimal interno - o Excel
    em portugues o desenha como virgula - e o '#' depois dele e placeholder
    OPCIONAL, que some quando nao ha digito. O separador nao some junto.

    Entao '#,##0.##' aplicado a 678 nao mostra '678': mostra '678,'. O valor
    gravado esta certo, a celula e numerica, e mesmo assim a coluna inteira sai
    com virgula sobrando na tela.

    A troca por FORMATO_INTEIRO quando o valor nao tem parte decimal e o que
    resolve, e por isso a mascara nunca e copiada direto da declaracao da coluna.
    """
    if formato != FORMATO_GERAL:
        return formato
    if isinstance(valor, bool) or not isinstance(valor, (int, float)):
        return formato
    return FORMATO_INTEIRO if float(valor).is_integer() else formato


def _situacao(dentro: bool | None) -> str:
    """Coluna de texto com o veredito da meta, para dar para filtrar na planilha."""
    if dentro is None:
        return "sem dado"
    return "na meta" if dentro else "fora da meta"


# ---------------------------------------------------------------------------
# Montagem de aba
# ---------------------------------------------------------------------------

def _escrever_aba(
    aba: Worksheet,
    colunas: list[tuple[str, int, str | None]],
    linhas: list[list],
    filtro: bool = True,
) -> None:
    """
    Escreve cabecalho formatado, dados e congela o painel.

    'colunas' e uma lista de (titulo, largura, formato). O formato None deixa a
    celula como texto ou como numero sem mascara.

    O painel congelado fica em A2: rolar 700 linhas de atraso sem o cabecalho a
    vista obriga a subir para lembrar o que e cada coluna.

    'filtro' desliga o filtro automatico. A aba Resumo tem linhas em branco
    separando secoes, e filtrar por cima delas esconderia blocos inteiros sem
    aviso - nas outras quatro, que sao tabelas de verdade, o filtro e o que
    torna a planilha util.
    """
    fundo = PatternFill("solid", fgColor=CABECALHO_FUNDO)
    fonte = Font(bold=True, color=CABECALHO_TEXTO)

    for indice, (titulo, largura, _) in enumerate(colunas, start=1):
        celula = aba.cell(row=1, column=indice, value=titulo)
        celula.fill = fundo
        celula.font = fonte
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        aba.column_dimensions[get_column_letter(indice)].width = largura

    for numero, linha in enumerate(linhas, start=2):
        for indice, (valor, (_, _, formato)) in enumerate(zip(linha, colunas), start=1):
            celula = aba.cell(row=numero, column=indice)
            # None de proposito: celula vazia, e nao o texto 'sem dado'. Uma
            # palavra no meio da coluna numerica quebraria MEDIA e ordenacao.
            celula.value = valor if not isinstance(valor, str) else _texto_seguro(valor)
            mascara = _mascara(formato, valor)
            if mascara and valor is not None:
                celula.number_format = mascara

    aba.freeze_panes = "A2"
    if filtro:
        aba.auto_filter.ref = (
            f"A1:{get_column_letter(len(colunas))}{max(len(linhas) + 1, 1)}"
        )


# ---------------------------------------------------------------------------
# As cinco abas
# ---------------------------------------------------------------------------

def _aba_resumo(
    aba: Worksheet, resultado: ResultadoIndicadores, cfg: Configuracao
) -> None:
    """KPIs consolidados, com valor, meta e situacao."""
    total = resultado.total
    dentro, medidas, pct_maquinas = maquinas_na_meta(resultado, cfg)
    meta_maquinas = cfg.metas["maquinas_na_meta_minimo_pct"]

    linhas = [
        [
            "Máquinas na meta de produtividade",
            pct_maquinas,
            meta_maquinas,
            "mínimo",
            _situacao(na_meta("maquinas_na_meta_pct", pct_maquinas, meta_maquinas)),
            f"{dentro} de {medidas} máquinas com produtividade medida no período",
        ]
    ]
    for rotulo, indicador, unidade in (
        ("Refugo", "refugo_pct", "%"),
        ("Tempo médio de acerto", "setup_medio_min", "min"),
        ("Entrega no prazo (OTD)", "otd_pct", "%"),
        ("Aproveitamento de tempo", "aproveitamento_pct", "%"),
    ):
        valor = getattr(total, indicador)
        meta = meta_de(indicador, cfg)
        sentido = "máximo" if indicador in ("refugo_pct", "setup_medio_min") else "mínimo"
        linhas.append([
            f"{rotulo} ({unidade})",
            valor,
            meta,
            sentido,
            _situacao(na_meta(indicador, valor, meta)),
            "",
        ])

    # As somas ficam na mesma aba dos indicadores de proposito: sao os numeros
    # que produzem cada um deles, e ter os dois lado a lado e o que permite
    # refazer a conta sem abrir outra aba.
    linhas += [
        [],
        ["Período por coorte de entrada", None, None, "", "", f"{resultado.de} a {resultado.ate}"],
        ["Ordens no período", total.ordens, None, "", "", ""],
        [],
        ["Somas do período", None, None, "", "", "de onde saem os indicadores acima"],
        ["Tiragem boa (folhas)", total.tiragem_boa, None, "", "", ""],
        ["Refugo (folhas)", total.refugo, None, "", "", ""],
        ["Tempo de produção (min)", total.tempo_producao_min, None, "", "", ""],
        ["Tempo de acerto (min)", total.tempo_setup_min, None, "", "", ""],
        ["Ordens com acerto preenchido", total.ordens_com_setup, None, "", "", ""],
        ["Entregues com prazo", total.entregues_com_prazo, None, "", "", ""],
        ["Entregues no prazo", total.entregues_no_prazo, None, "", "", ""],
        [],
        [
            "Célula vazia na coluna Valor",
            None, None, "", "",
            "indicador sem dado: não foi medido no período, e por isso não está "
            "nem na meta nem fora dela",
        ],
    ]
    if total.entregues_sem_prazo:
        linhas.append([
            "Entregues sem prazo",
            total.entregues_sem_prazo, None, "", "",
            "ficaram fora do OTD: sem prazo não há como dizer se atrasou",
        ])
    if resultado.mes_incompleto:
        linhas.append([
            "Período incompleto",
            None, None, "", "",
            f"{resultado.mes_incompleto} sofre censura à direita: a exportação foi "
            f"cortada dentro dele e o OTD do mês é otimista",
        ])

    _escrever_aba(
        aba,
        [
            ("Indicador", 32, None),
            ("Valor", 14, FORMATO_GERAL),
            ("Meta", 12, FORMATO_GERAL),
            ("Sentido", 10, None),
            ("Situação", 14, None),
            ("Observação", 62, None),
        ],
        linhas,
        filtro=False,
    )


def _aba_por_maquina(
    aba: Worksheet, resultado: ResultadoIndicadores, cfg: Configuracao
) -> None:
    """Indicadores por equipamento, com a meta de produtividade de cada um."""
    linhas = []
    for nome, ind in resultado.por_maquina.items():
        linhas.append([
            nome,
            ind.ordens,
            ind.tiragem_boa,
            ind.refugo,
            ind.refugo_pct,
            ind.tempo_producao_min,
            ind.tempo_setup_min,
            ind.produtividade,
            cfg.metas_produtividade.get(nome),
            _situacao(produtividade_na_meta(nome, ind, cfg)),
            ind.setup_medio_min,
            ind.otd_pct,
            ind.aproveitamento_pct,
        ])

    _escrever_aba(
        aba,
        [
            ("Máquina", 24, None),
            ("Ordens", 9, FORMATO_INTEIRO),
            ("Tiragem boa", 14, FORMATO_INTEIRO),
            ("Refugo", 12, FORMATO_INTEIRO),
            ("Refugo %", 11, FORMATO_PERCENTUAL),
            ("Produção (min)", 14, FORMATO_UM_DECIMAL),
            ("Acerto (min)", 13, FORMATO_UM_DECIMAL),
            ("Folhas/hora", 13, FORMATO_INTEIRO),
            ("Meta folhas/hora", 16, FORMATO_INTEIRO),
            ("Situação", 14, None),
            ("Acerto médio (min)", 17, FORMATO_UM_DECIMAL),
            ("OTD %", 10, FORMATO_UM_DECIMAL),
            ("Aproveitamento %", 17, FORMATO_UM_DECIMAL),
        ],
        linhas,
    )


def _aba_por_cliente(aba: Worksheet, por_cliente: dict[str, Indicadores]) -> None:
    """
    Volume e refugo por cliente, do maior volume para o menor.

    Sem produtividade nem acerto: esses sao propriedades da maquina, e
    atribui-los ao cliente sugeriria que o cliente e responsavel por eles.
    """
    ordenado = sorted(
        por_cliente.items(), key=lambda item: item[1].tiragem_boa, reverse=True
    )
    linhas = [
        [
            cliente,
            ind.ordens,
            ind.tiragem_boa,
            ind.refugo,
            ind.refugo_pct,
            ind.entregues_com_prazo,
            ind.otd_pct,
        ]
        for cliente, ind in ordenado
    ]

    _escrever_aba(
        aba,
        [
            ("Cliente", 16, None),
            ("Ordens", 9, FORMATO_INTEIRO),
            ("Tiragem boa", 14, FORMATO_INTEIRO),
            ("Refugo", 12, FORMATO_INTEIRO),
            ("Refugo %", 11, FORMATO_PERCENTUAL),
            ("Entregues com prazo", 19, FORMATO_INTEIRO),
            ("OTD %", 10, FORMATO_UM_DECIMAL),
        ],
        linhas,
    )


def _aba_atrasos(aba: Worksheet, atrasos: list) -> None:
    """
    Todas as ordens entregues fora do prazo.

    A planilha leva a lista inteira, e nao as dez do painel: o painel e para
    olhar de longe, e aqui alguem vai filtrar por cliente ou por maquina para
    descobrir onde o atraso se concentra.
    """
    linhas = [
        [
            atraso.numero_os,
            atraso.cliente,
            atraso.produto,
            atraso.maquina,
            date.fromisoformat(atraso.data_prazo),
            date.fromisoformat(atraso.data_entrega),
            atraso.dias,
        ]
        for atraso in atrasos
    ]
    _escrever_aba(
        aba,
        [
            ("OS", 16, None),
            ("Cliente", 12, None),
            ("Produto", 40, None),
            ("Máquina", 22, None),
            ("Prazo", 12, FORMATO_DATA),
            ("Entrega", 12, FORMATO_DATA),
            ("Atraso (dias)", 14, FORMATO_INTEIRO),
        ],
        linhas,
    )


def _aba_qualidade(aba: Worksheet, ocorrencias: list[Ocorrencia]) -> None:
    """
    O que foi corrigido, rejeitado, descartado e deduplicado na ingestao.

    Mesmas colunas de 'saida/relatorio_qualidade.csv' da secao 6.3, para os dois
    arquivos poderem ser conferidos um contra o outro linha a linha.

    Nao e recortado pelo periodo: a ingestao acontece sobre a pasta inteira, e
    filtrar por data esconderia justamente a linha rejeitada por ter data
    ilegivel.
    """
    linhas = [
        [
            ocorrencia.arquivo,
            ocorrencia.linha_planilha,
            ocorrencia.coluna,
            ocorrencia.valor_original,
            ocorrencia.problema,
            ocorrencia.acao,
            ocorrencia.valor_final,
        ]
        for ocorrencia in ocorrencias
    ]
    _escrever_aba(
        aba,
        [
            ("Arquivo", 34, None),
            ("Linha", 8, FORMATO_INTEIRO),
            ("Coluna", 18, None),
            ("Valor original", 26, None),
            ("Problema", 70, None),
            ("Ação", 12, None),
            ("Valor final", 24, None),
        ],
        linhas,
    )


# ---------------------------------------------------------------------------
# Montagem do arquivo
# ---------------------------------------------------------------------------

def montar_livro(
    resultado: ResultadoIndicadores,
    por_cliente: dict[str, Indicadores],
    atrasos: list,
    ocorrencias: list[Ocorrencia],
    cfg: Configuracao,
) -> Workbook:
    """Monta o Workbook com as cinco abas da secao 6.2, na ordem declarada."""
    resumo_nome, maquina_nome, cliente_nome, atrasos_nome, qualidade_nome = (
        ABAS_ESPERADAS
    )

    livro = Workbook()
    # O Workbook nasce com uma aba; renomeia-la em vez de criar outra evita
    # deixar uma 'Sheet' vazia no arquivo final.
    resumo = livro.active
    resumo.title = resumo_nome

    _aba_resumo(resumo, resultado, cfg)
    _aba_por_maquina(livro.create_sheet(maquina_nome), resultado, cfg)
    _aba_por_cliente(livro.create_sheet(cliente_nome), por_cliente)
    _aba_atrasos(livro.create_sheet(atrasos_nome), atrasos)
    _aba_qualidade(livro.create_sheet(qualidade_nome), ocorrencias)
    return livro


def gerar(
    cfg: Configuracao,
    banco: Path | None = None,
    entrada: Path | None = None,
    de: str | None = None,
    ate: str | None = None,
    destino: Path | None = None,
    ocorrencias: list[Ocorrencia] | None = None,
) -> tuple[Path, int]:
    """
    Gera o relatorio e devolve o caminho e o numero de ordens no periodo.

    'ocorrencias' vem de fora quando quem chama ja importou - e o caso do
    'tudo' do Bloco 10, que nao pode importar duas vezes so para preencher uma
    aba. Omitido, o modulo roda a importacao ele mesmo: ela e idempotente, entao
    o banco nao muda, e e a unica forma de a aba de qualidade existir sem uma
    tabela de ocorrencias no banco.
    """
    if ocorrencias is None:
        importacao = importar(cfg, entrada=entrada, banco=banco)
        ocorrencias = importacao.ocorrencias

    caminho_banco = cfg.caminho("banco") if banco is None else Path(banco)
    if not caminho_banco.is_file():
        raise ErroRelatorio(
            f"banco nao encontrado em {caminho_banco}. "
            f"Rode 'python -m src.banco' antes"
        )

    destino = (
        cfg.caminho("saida") / ARQUIVO_PADRAO if destino is None else Path(destino)
    )

    conexao = conectar(caminho_banco)
    try:
        resultado = calcular(conexao, de, ate)
        por_cliente = calcular_corte(conexao, "cliente", resultado.de, resultado.ate)
        atrasos = ordens_atrasadas(conexao, resultado, limite=None)
    finally:
        conexao.close()

    livro = montar_livro(resultado, por_cliente, atrasos, ocorrencias, cfg)

    try:
        destino.parent.mkdir(parents=True, exist_ok=True)
        livro.save(destino)
    except OSError as erro:
        raise ErroRelatorio(f"nao foi possivel gravar {destino}: {erro}") from erro

    return destino, resultado.total.ordens


def _argumentos(argv: list[str]) -> argparse.Namespace:
    analisador = argparse.ArgumentParser(
        prog="python -m src.relatorio_excel",
        description="Gera o relatorio Excel formatado com as cinco abas.",
    )
    analisador.add_argument(
        "--config", default=str(CONFIG_PADRAO), help="caminho do config.yaml"
    )
    analisador.add_argument("--banco", help="sobrepoe caminhos.banco")
    analisador.add_argument("--entrada", help="sobrepoe caminhos.entrada")
    analisador.add_argument("--saida", help="sobrepoe o arquivo de destino")
    analisador.add_argument("--de", help="inicio do periodo, AAAA-MM-DD")
    analisador.add_argument("--ate", help="fim do periodo, AAAA-MM-DD")
    return analisador.parse_args(argv)


def main(argv: list[str]) -> int:
    """
    Entrada de 'python -m src.relatorio_excel'.

    Importa antes de montar, porque a aba de qualidade precisa das ocorrencias
    da ingestao, que o banco nao guarda linha a linha. A importacao e
    idempotente: rodar de novo nao duplica ordem nenhuma.
    """
    args = _argumentos(argv)

    try:
        cfg = carregar(args.config)
    except ErroConfiguracao as erro:
        print(f"ERRO DE CONFIGURACAO: {erro}", file=sys.stderr)
        return CODIGO_SAIDA_ERRO_CONFIG

    try:
        destino, ordens = gerar(
            cfg,
            banco=args.banco,
            entrada=args.entrada,
            de=args.de,
            ate=args.ate,
            destino=args.saida,
        )
    except (ErroRelatorio, ErroBanco, ErroIndicadores, ErroLeitura, ErroValidacao) as erro:
        print(f"ERRO DE RELATORIO: {erro}", file=sys.stderr)
        return 1

    print(f"Relatorio gravado em {destino}")
    print(f"{ordens} ordens no periodo, em 5 abas")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
