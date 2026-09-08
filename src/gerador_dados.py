"""
Gerador de planilhas de producao sinteticas.

Produz arquivos .xlsx com a cara do que uma fabrica realmente exporta: cabecalho
fora da primeira linha, rodape com linha TOTAL, nome de maquina digitado de
qualquer jeito, data em quatro formatos, numero como texto, celula vazia, linha
duplicada, status incoerente e valor impossivel.

Modelo plano, uma linha = uma ordem, so a etapa de impressao. A secao 4.4 da
especificacao explica o que essa simplificacao custa.

Nenhum dado de empresa real. Cliente e operador sao codigo ('CLI-001', 'OP-03').

A semente vem do config.yaml: mesma semente, mesmo conteudo gerado.
"""

import argparse
import random
import sys
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.config import (
    CODIGO_SAIDA_ERRO_CONFIG,
    CONFIG_PADRAO,
    Configuracao,
    ErroConfiguracao,
    carregar,
)
# A base do serial, a tabela de meses e os nomes canonicos de status moram na
# limpeza, que e quem precisa le-los de volta. O gerador so escreve, entao
# importa de la em vez de manter uma segunda copia que sairia de sincronia na
# primeira correcao. Os tres status precisam existir na secao 'status' do
# config.yaml - gerar() confere antes de comecar.
from src.limpeza import (
    BASE_SERIAL_EXCEL,
    MESES_ABREV,
    STATUS_CANCELADA,
    STATUS_EM_PRODUCAO,
    STATUS_ENTREGUE,
)

# Fim da janela de historico. Fixo de proposito: ancorar em date.today() faria a
# mesma semente produzir datas diferentes a cada dia. Pode ser trocado por --fim.
DATA_FIM_PADRAO = date(2026, 6, 30)

# Botoes de realismo que nao viraram config por serem do gerador, nao do chao de
# fabrica. Velocidade, setup, refugo, tiragem e atraso todos saem do config.yaml.
FAIXA_FATOR_VELOCIDADE = (0.80, 1.20)   # sobre metas_produtividade da maquina
FAIXA_RUIDO_JOB = (0.85, 1.15)          # variacao de uma ordem para outra
FAIXA_DIAS_PRAZO = (3, 15)              # prazo prometido, em dias corridos

SUJEIRA_MAQUINA = "nome de maquina em variante"
SUJEIRA_DATA = "data em formato alternativo"
SUJEIRA_NUMERO_TEXTO = "numero como texto"
SUJEIRA_CELULA_VAZIA = "celula vazia"
SUJEIRA_DUPLICADA = "linha duplicada"
SUJEIRA_IMPOSSIVEL = "valor impossivel"
SUJEIRA_ESPACO = "espaco sobrando"
SUJEIRA_STATUS_VARIANTE = "status em variante"
SUJEIRA_STATUS_INCOERENTE = "status incoerente com a entrega"
SUJEIRA_CABECALHO = "cabecalho fora da linha 1"
SUJEIRA_RODAPE = "lixo no rodape"

SUJEIRAS_POR_LINHA = (
    SUJEIRA_MAQUINA,
    SUJEIRA_DATA,
    SUJEIRA_NUMERO_TEXTO,
    SUJEIRA_CELULA_VAZIA,
    SUJEIRA_DUPLICADA,
    SUJEIRA_IMPOSSIVEL,
    SUJEIRA_ESPACO,
    SUJEIRA_STATUS_VARIANTE,
    SUJEIRA_STATUS_INCOERENTE,
)

CAMPOS_TEXTO = (
    "numero_os", "cliente", "produto", "maquina", "operador", "turno", "status",
)
CAMPOS_DATA = ("data_entrada", "data_prazo", "data_entrega")
CAMPOS_NUMERO = (
    "tiragem_prevista",
    "tiragem_boa",
    "refugo",
    "tempo_setup_min",
    "tempo_producao_min",
)
# Somados na linha TOTAL do rodape, como faria um operador no Excel.
CAMPOS_SOMADOS = CAMPOS_NUMERO


class ErroGeracao(Exception):
    """Parametro de geracao incoerente ou pasta de destino inacessivel."""


@dataclass
class Ordem:
    """Uma ordem de producao limpa, antes de qualquer sujeira ser aplicada."""

    numero_os: str
    cliente: str
    produto: str
    maquina: str
    operador: str
    turno: str
    status: str
    data_entrada: date
    data_prazo: date
    data_entrega: date | None
    tiragem_prevista: int
    tiragem_boa: int
    refugo: int
    tempo_setup_min: float
    tempo_producao_min: float


@dataclass
class ResumoGeracao:
    """O que a geracao produziu, para conferencia e para o log."""

    arquivos: list[Path] = field(default_factory=list)
    linhas_por_arquivo: dict[str, int] = field(default_factory=dict)
    ordens_limpas: int = 0
    sujeiras: dict[str, int] = field(default_factory=dict)
    sujeiras_sem_efeito: dict[str, int] = field(default_factory=dict)
    arquivos_preexistentes: list[Path] = field(default_factory=list)
    por_status: dict[str, int] = field(default_factory=dict)
    otd_por_mes: dict[str, tuple[int, int]] = field(default_factory=dict)
    refugo_por_maquina: dict[str, tuple[int, int]] = field(default_factory=dict)
    probabilidade_base_otd: float = 0.0

    def contar(self, sujeira: str) -> None:
        self.sujeiras[sujeira] = self.sujeiras.get(sujeira, 0) + 1

    def contar_sem_efeito(self, sujeira: str) -> None:
        """
        Registra a sujeira sorteada que nao mudou a celula.

        Nao vira sorteio novo de proposito: repetir o sorteio consumiria numeros
        diferentes do gerador e a mesma semente passaria a produzir outro
        historico, invalidando os numeros de OTD e refugo ja conferidos. A
        tentativa perdida fica visivel aqui em vez de inflar 'sujeiras'.
        """
        self.sujeiras_sem_efeito[sujeira] = self.sujeiras_sem_efeito.get(sujeira, 0) + 1


# ---------------------------------------------------------------------------
# Formatacao de valores sujos
# ---------------------------------------------------------------------------

def para_serial_excel(valor: date) -> int:
    """Converte uma data no numero serial do Excel (sistema de datas 1900)."""
    return (valor - BASE_SERIAL_EXCEL).days


def formatar_data_suja(valor: date, rng: random.Random) -> str | int:
    """
    Devolve a data num dos quatro formatos da secao 4.2.

    O serial sai como numero inteiro, nao como texto: e assim que ele aparece
    quando uma coluna de data perde a formatacao no Excel.
    """
    formato = rng.choice(("barra", "iso", "abreviado", "serial"))
    if formato == "barra":
        return f"{valor.day:02d}/{valor.month:02d}/{valor.year}"
    if formato == "iso":
        return valor.isoformat()
    if formato == "abreviado":
        return f"{valor.day}-{MESES_ABREV[valor.month - 1]}-{valor.year % 100:02d}"
    return para_serial_excel(valor)


def formatar_numero_sujo(valor: float, rng: random.Random) -> str:
    """Devolve o numero como texto: milhar com ponto, com espaco, ou decimal com virgula."""
    inteiro = int(round(valor))
    formato = rng.choice(("milhar_ponto", "milhar_espaco", "virgula_decimal"))
    if formato == "milhar_ponto":
        return f"{inteiro:,}".replace(",", ".")
    if formato == "milhar_espaco":
        return f"{inteiro:,}".replace(",", " ")
    return f"{valor:.1f}".replace(".", ",")


def sujar_espacos(texto: str, rng: random.Random) -> str:
    """Acrescenta espaco na ponta e, as vezes, no meio do texto."""
    sujo = texto
    if " " in sujo and rng.random() < 0.5:
        partes = sujo.split(" ", 1)
        sujo = f"{partes[0]}  {partes[1]}"
    return f"{' ' * rng.randint(0, 2)}{sujo}{' ' * rng.randint(1, 3)}"


def _variante_com_caixa(variantes: list[str], rng: random.Random) -> str:
    """
    Sorteia uma variante do config e mexe na caixa.

    As variantes no config estao todas em minuscula. A variacao de caixa e
    aplicada aqui para cobrir os exemplos da secao 4.2 ('HEIDELBERG  SM 52 ').
    """
    escolhida = rng.choice(variantes)
    caixa = rng.choice(("baixa", "alta", "titulo"))
    if caixa == "alta":
        escolhida = escolhida.upper()
    elif caixa == "titulo":
        escolhida = escolhida.title()
    if rng.random() < 0.4:
        escolhida = escolhida.replace(" ", "  ", 1)
    return escolhida


# ---------------------------------------------------------------------------
# Geracao das ordens limpas
# ---------------------------------------------------------------------------

def _meses_da_janela(data_fim: date, meses: int) -> list[tuple[int, int]]:
    """Devolve os (ano, mes) da janela, do mais antigo para o mais recente."""
    janela = []
    ano, mes = data_fim.year, data_fim.month
    for _ in range(meses):
        janela.append((ano, mes))
        mes -= 1
        if mes == 0:
            ano, mes = ano - 1, 12
    return list(reversed(janela))


def _primeiro_e_ultimo_dia(ano: int, mes: int, data_fim: date) -> tuple[date, date]:
    """Limites do mes, truncados pelo fim da janela."""
    primeiro = date(ano, mes, 1)
    if mes == 12:
        ultimo = date(ano, 12, 31)
    else:
        ultimo = date(ano, mes + 1, 1) - timedelta(days=1)
    return primeiro, min(ultimo, data_fim)


def _dia_util(inicio: date, fim: date, rng: random.Random) -> date:
    """
    Sorteia um dia entre inicio e fim, evitando sabado e domingo.

    A fabrica tambem roda no fim de semana de vez em quando, entao o desvio
    para segunda nao e absoluto.
    """
    dia = inicio + timedelta(days=rng.randint(0, (fim - inicio).days))
    if dia.weekday() >= 5 and rng.random() < 0.85:
        dia += timedelta(days=7 - dia.weekday())
        if dia > fim:
            dia = fim
    return dia


def _velocidades(cfg: Configuracao, rng: random.Random) -> dict[str, float]:
    """
    Sorteia a velocidade real de cada maquina uma vez so, em volta da meta.

    Assim uma maquina e consistentemente mais rapida que outra ao longo de todo
    o historico, em vez de oscilar. A ancora e 'metas_produtividade': nenhum
    numero por maquina fica cravado aqui.
    """
    return {
        maquina: meta * rng.uniform(*FAIXA_FATOR_VELOCIDADE)
        for maquina, meta in cfg.metas_produtividade.items()
    }


def _escolher_maquina(tiragem: int, cfg: Configuracao, rng: random.Random) -> str:
    """
    Roteia a ordem entre digital e offset conforme a tiragem.

    Tiragem curta pode ir para qualquer uma, com peso maior para a digital, que
    e onde ela e economica. Acima do limite so sobra offset.
    """
    digitais = cfg.gerador["maquinas_digitais"]
    offset = [m for m in cfg.maquinas if m not in digitais]

    if tiragem > cfg.gerador["tiragem_maxima_digital"]:
        return rng.choice(offset)

    peso = cfg.gerador["peso_digital_tiragem_curta"]
    return rng.choice(offset + list(digitais) * peso)


def _definir_status(
    data_entrada: date,
    data_fim: date,
    cfg: Configuracao,
    rng: random.Random,
) -> str:
    """Sorteia a situacao da ordem no momento em que a planilha foi exportada."""
    parametros = cfg.gerador["status"]

    if rng.random() < parametros["prob_cancelada"]:
        return STATUS_CANCELADA

    dias_ate_o_fim = (data_fim - data_entrada).days
    if (
        dias_ate_o_fim <= parametros["dias_janela_em_producao"]
        and rng.random() < parametros["prob_em_producao"]
    ):
        return STATUS_EM_PRODUCAO

    return STATUS_ENTREGUE


def _montar_ordem(
    sequencial: int,
    data_entrada: date,
    data_fim: date,
    cfg: Configuracao,
    velocidades: dict[str, float],
    rng: random.Random,
) -> Ordem:
    """
    Monta uma ordem coerente, ainda sem data de entrega.

    A entrega e atribuida depois, em _atribuir_entregas, porque a probabilidade
    de atraso depende de quantas ordens entregues existem no total.
    """
    produto = rng.choice(cfg.gerador["produtos"])
    minimo, maximo = cfg.gerador["tiragem_por_produto"][produto]
    # Tiragem arredondada na centena: ninguem pede 4.317 folhas.
    tiragem_prevista = max(100, int(round(rng.uniform(minimo, maximo), -2)))

    maquina = _escolher_maquina(tiragem_prevista, cfg, rng)
    status = _definir_status(data_entrada, data_fim, cfg, rng)

    setup_minimo, setup_maximo = cfg.gerador["setup_por_maquina"][maquina]
    tempo_setup_min = round(rng.uniform(setup_minimo, setup_maximo), 1)

    if status == STATUS_CANCELADA:
        # O cliente cancelou depois de a maquina ser acertada: o setup foi
        # gasto, nada foi produzido. E daqui que sai a ordem com
        # tempo_producao_min = 0 que a guarda de divisao por zero da secao 5
        # precisa encontrar em dado real.
        tiragem_boa, refugo, tempo_producao_min = 0, 0, 0.0
    else:
        # Um sorteio so decide as tres fatias do fechamento. 'Fechar exato' e a
        # tiragem boa bater com o pedido - o que a grafica entrega ao cliente,
        # nao a soma contabil boa + refugo.
        fechamento = cfg.gerador["fechamento_tiragem"]
        sorteio = rng.random()
        if sorteio < fechamento["prob_exato"]:
            margem_pct = 0.0
        elif sorteio < fechamento["prob_exato"] + fechamento["prob_abaixo_do_pedido"]:
            margem_pct = -rng.uniform(0, fechamento["falta_maxima_pct"])
        else:
            # Roda-se folha a mais para cobrir o acerto, entao o normal e
            # fechar acima do pedido.
            margem_pct = rng.uniform(0, fechamento["sobra_maxima_pct"])
        tiragem_boa = max(1, int(round(tiragem_prevista * (1 + margem_pct / 100))))

        # refugo_pct e definido como refugo/(boa+refugo), igual a secao 5.
        # Invertendo a formula, o refugo que produz esse percentual e
        # boa * pct / (100 - pct).
        refugo_pct = rng.uniform(*cfg.gerador["refugo_por_maquina"][maquina])
        refugo = int(round(tiragem_boa * refugo_pct / (100 - refugo_pct)))

        # O ruido por job e o que faz a produtividade variar em volta da
        # velocidade da maquina. Sem ele, toda ordem daria o mesmo folhas/hora.
        ruido = rng.uniform(*FAIXA_RUIDO_JOB)
        tempo_producao_min = round(tiragem_boa / (velocidades[maquina] * ruido) * 60, 1)

    return Ordem(
        numero_os=f"OS-{data_entrada.year}-{sequencial:04d}",
        cliente=rng.choice(cfg.gerador["clientes"]),
        produto=produto,
        maquina=maquina,
        operador=rng.choice(cfg.gerador["operadores"]),
        turno=rng.choice(list(cfg.turnos)),
        status=status,
        data_entrada=data_entrada,
        data_prazo=data_entrada + timedelta(days=rng.randint(*FAIXA_DIAS_PRAZO)),
        data_entrega=None,
        tiragem_prevista=tiragem_prevista,
        tiragem_boa=tiragem_boa,
        refugo=refugo,
        tempo_setup_min=tempo_setup_min,
        tempo_producao_min=tempo_producao_min,
    )


def _na_concentracao(ordem: Ordem, cfg: Configuracao) -> bool:
    """
    Diz se a ordem cai no periodo e na maquina em crise.

    O agrupamento e pelo mes de ENTRADA, nao pelo mes de entrega. E preciso que
    o Bloco 7 agrupe o OTD mensal do mesmo jeito, senao a queda aparece
    borrada entre dois meses em vez de concentrada em um.
    """
    concentracao = cfg.gerador["atrasos"]["concentracao"]
    return (
        ordem.maquina == concentracao["maquina"]
        and f"{ordem.data_entrada:%Y-%m}" == concentracao["mes"]
    )


def _probabilidade_base_otd(ordens: list[Ordem], cfg: Configuracao) -> float:
    """
    Resolve a probabilidade de entrega no prazo fora do periodo em crise.

    O config declara o OTD alvo do historico inteiro e o OTD da concentracao.
    A probabilidade do resto sai por diferenca, para que a media ponderada bata
    com o alvo em vez de ficar abaixo dele sem ninguem notar.
    """
    atrasos = cfg.gerador["atrasos"]
    alvo = atrasos["otd_alvo_pct"] / 100
    otd_concentracao = atrasos["concentracao"]["otd_pct"] / 100

    entregues = [o for o in ordens if o.status == STATUS_ENTREGUE]
    na_crise = sum(1 for o in entregues if _na_concentracao(o, cfg))
    fora = len(entregues) - na_crise

    if fora == 0:
        raise ErroGeracao(
            f"todas as {len(entregues)} ordens entregues caem na concentracao "
            f"({atrasos['concentracao']['maquina']} em "
            f"{atrasos['concentracao']['mes']}), entao nao ha como atingir o OTD "
            f"alvo de {atrasos['otd_alvo_pct']}% fora dela"
        )

    base = (alvo * len(entregues) - na_crise * otd_concentracao) / fora

    if not 0 <= base <= 1:
        raise ErroGeracao(
            f"nao da para atingir 'otd_alvo_pct' de {atrasos['otd_alvo_pct']}% com "
            f"{na_crise} de {len(entregues)} ordens entregues concentradas em "
            f"{otd_concentracao:.0%}: as demais precisariam de "
            f"{base:.1%} de entrega no prazo. Aproxime 'otd_alvo_pct' de "
            f"'concentracao.otd_pct' ou reduza o tamanho da concentracao"
        )

    return base


def _atribuir_entregas(
    ordens: list[Ordem],
    cfg: Configuracao,
    data_fim: date,
    rng: random.Random,
) -> float:
    """
    Preenche a data de entrega das ordens entregues e devolve a base usada.

    Ordem cuja entrega cairia depois do fim da janela vira EM PRODUCAO: uma
    planilha exportada em 30/06 nao pode conter entrega em 10/07. Isso tira do
    denominador do OTD mais ordens atrasadas que pontuais, entao o OTD apurado
    fica um pouco acima do alvo - o resumo mostra o valor real.
    """
    base = _probabilidade_base_otd(ordens, cfg)
    minimo_atraso, maximo_atraso = cfg.gerador["atrasos"]["dias_atraso"]

    for ordem in ordens:
        if ordem.status != STATUS_ENTREGUE:
            continue

        if _na_concentracao(ordem, cfg):
            prazo_cumprido = rng.random() < cfg.gerador["atrasos"]["concentracao"]["otd_pct"] / 100
        else:
            prazo_cumprido = rng.random() < base

        if prazo_cumprido:
            dias = (ordem.data_prazo - ordem.data_entrada).days
            entrega = ordem.data_entrada + timedelta(days=rng.randint(1, dias))
        else:
            entrega = ordem.data_prazo + timedelta(
                days=rng.randint(minimo_atraso, maximo_atraso)
            )

        if entrega > data_fim:
            ordem.status = STATUS_EM_PRODUCAO
        else:
            ordem.data_entrega = entrega

    return base


def gerar_ordens(
    cfg: Configuracao,
    meses: int,
    data_fim: date,
    rng: random.Random,
) -> tuple[list[Ordem], float]:
    """Gera as ordens limpas de toda a janela, em ordem cronologica."""
    velocidades = _velocidades(cfg, rng)
    ordens_por_mes = cfg.gerador["ordens_por_mes"]

    ordens: list[Ordem] = []
    sequencial = 0
    for ano, mes in _meses_da_janela(data_fim, meses):
        primeiro, ultimo = _primeiro_e_ultimo_dia(ano, mes, data_fim)
        do_mes = []
        for _ in range(ordens_por_mes):
            sequencial += 1
            entrada = _dia_util(primeiro, ultimo, rng)
            do_mes.append(
                _montar_ordem(sequencial, entrada, data_fim, cfg, velocidades, rng)
            )
        do_mes.sort(key=lambda o: o.data_entrada)
        ordens.extend(do_mes)

    base = _atribuir_entregas(ordens, cfg, data_fim, rng)
    return ordens, base


# ---------------------------------------------------------------------------
# Injecao de sujeira
# ---------------------------------------------------------------------------

def celulas_limpas(ordem: Ordem, cfg: Configuracao) -> dict[str, Any]:
    """Converte a ordem no dicionario de celulas, na ordem das colunas do config."""
    valores = {
        "numero_os": ordem.numero_os,
        "cliente": ordem.cliente,
        "produto": ordem.produto,
        "maquina": ordem.maquina,
        "operador": ordem.operador,
        "turno": ordem.turno,
        "status": ordem.status,
        "data_entrada": ordem.data_entrada,
        "data_prazo": ordem.data_prazo,
        "data_entrega": ordem.data_entrega,
        "tiragem_prevista": ordem.tiragem_prevista,
        "tiragem_boa": ordem.tiragem_boa,
        "refugo": ordem.refugo,
        "tempo_setup_min": ordem.tempo_setup_min,
        "tempo_producao_min": ordem.tempo_producao_min,
    }
    faltando = set(cfg.colunas) - set(valores)
    if faltando:
        raise ErroGeracao(
            f"o config.yaml declara colunas que o gerador nao sabe preencher: "
            f"{', '.join(sorted(faltando))}. Colunas geradas: "
            f"{', '.join(sorted(valores))}"
        )
    return {campo: valores[campo] for campo in cfg.colunas}


def _sujar_maquina(celulas: dict, cfg: Configuracao, rng: random.Random) -> None:
    celulas["maquina"] = _variante_com_caixa(cfg.maquinas[celulas["maquina"]], rng)


def _sujar_status(celulas: dict, cfg: Configuracao, rng: random.Random) -> None:
    celulas["status"] = _variante_com_caixa(cfg.status[celulas["status"]], rng)


def _sujar_data(celulas: dict, rng: random.Random) -> None:
    """Troca o formato de uma das datas preenchidas da linha."""
    preenchidas = [c for c in CAMPOS_DATA if celulas[c] is not None]
    campo = rng.choice(preenchidas)
    celulas[campo] = formatar_data_suja(celulas[campo], rng)


def _sujar_numero(celulas: dict, rng: random.Random) -> None:
    campo = rng.choice(CAMPOS_NUMERO)
    celulas[campo] = formatar_numero_sujo(celulas[campo], rng)


def _sujar_celula_vazia(celulas: dict, cfg: Configuracao, rng: random.Random) -> None:
    """
    Esvazia um campo obrigatorio.

    Os campos vem de 'validacao.obrigatorios' no config, mais tiragem_boa, que e
    obrigatoria so quando o status e ENTREGUE. Cravar a lista aqui faria mudar o
    config deixar a sujeira sem efeito, sem ninguem perceber.

    DATA ENTREGA ficou fora de proposito: com a coluna STATUS, a ausencia de
    data de entrega passou a ter um significado so - a ordem nao foi entregue.
    Esvaziar aleatoriamente era a unica sujeira que o pipeline nao tinha como
    detectar, porque virava indistinguivel de uma ordem legitimamente aberta.
    """
    candidatos = list(cfg.validacao["obrigatorios"]) + ["tiragem_boa"]
    celulas[rng.choice(candidatos)] = None


def _sujar_status_incoerente(celulas: dict, rng: random.Random) -> None:
    """
    Poe status e data de entrega em contradicao.

    Detectavel: ENTREGUE sem data de entrega, ou nao entregue com data. E a
    sujeira que substituiu o esvaziamento cego da data de entrega.
    """
    if celulas["status"] == STATUS_ENTREGUE:
        celulas["data_entrega"] = None
    else:
        celulas["data_entrega"] = celulas["data_entrada"] + timedelta(
            days=rng.randint(1, 20)
        )


def _sujar_valor_impossivel(celulas: dict, rng: random.Random) -> None:
    """Refugo demais, tempo negativo ou entrega antes da entrada."""
    possiveis = ["refugo", "tempo_negativo"]
    if celulas["data_entrega"] is not None:
        possiveis.append("entrega_antes")

    tipo = rng.choice(possiveis)
    if tipo == "refugo":
        # Acima de metade da tiragem, para estourar refugo_maximo_sobre_tiragem.
        celulas["refugo"] = int(celulas["tiragem_prevista"] * rng.uniform(0.6, 1.4))
    elif tipo == "tempo_negativo":
        campo = rng.choice(("tempo_setup_min", "tempo_producao_min"))
        celulas[campo] = -abs(celulas[campo])
    else:
        celulas["data_entrega"] = celulas["data_entrada"] - timedelta(
            days=rng.randint(1, 20)
        )


def _sujar_espacos(celulas: dict, rng: random.Random) -> None:
    """Espaco sobrando em alguns campos de texto da linha."""
    quantos = rng.randint(1, len(CAMPOS_TEXTO))
    for campo in rng.sample(CAMPOS_TEXTO, quantos):
        celulas[campo] = sujar_espacos(str(celulas[campo]), rng)


def sujar_linhas(
    ordens: list[Ordem],
    cfg: Configuracao,
    rng: random.Random,
    resumo: ResumoGeracao,
) -> list[dict[str, Any]]:
    """
    Converte as ordens em linhas de planilha e aplica a sujeira por linha.

    Cada linha sorteada recebe um tipo de sujeira. Uma linha so recebe uma:
    assim a contagem por tipo no resumo corresponde ao que esta na planilha, e
    da para conferir a olho.

    A contagem e feita DEPOIS de escrever, comparando a celula com o valor
    original. Nem toda sujeira sorteada suja: 'entregue'.upper() e o proprio
    nome canonico 'ENTREGUE', e negar o tempo de uma ordem cancelada da -0.0,
    que a planilha grava como zero. Contar no sorteio fazia o resumo prometer
    sujeira que nao esta no arquivo, e o pipeline levava a culpa por nao
    encontra-la.
    """
    taxa = cfg.gerador["taxa_sujeira"]
    linhas: list[dict[str, Any]] = []

    for ordem in ordens:
        celulas = celulas_limpas(ordem, cfg)

        if rng.random() >= taxa:
            linhas.append(celulas)
            continue

        antes = dict(celulas)
        sujeira = rng.choice(SUJEIRAS_POR_LINHA)
        if sujeira == SUJEIRA_MAQUINA:
            _sujar_maquina(celulas, cfg, rng)
        elif sujeira == SUJEIRA_DATA:
            _sujar_data(celulas, rng)
        elif sujeira == SUJEIRA_NUMERO_TEXTO:
            _sujar_numero(celulas, rng)
        elif sujeira == SUJEIRA_CELULA_VAZIA:
            _sujar_celula_vazia(celulas, cfg, rng)
        elif sujeira == SUJEIRA_IMPOSSIVEL:
            _sujar_valor_impossivel(celulas, rng)
        elif sujeira == SUJEIRA_ESPACO:
            _sujar_espacos(celulas, rng)
        elif sujeira == SUJEIRA_STATUS_VARIANTE:
            _sujar_status(celulas, cfg, rng)
        elif sujeira == SUJEIRA_STATUS_INCOERENTE:
            _sujar_status_incoerente(celulas, rng)

        linhas.append(celulas)

        if sujeira == SUJEIRA_DUPLICADA:
            # A mesma OS apontada duas vezes, uma logo abaixo da outra - o que
            # acontece quando o operador clica salvar duas vezes. Copia exata:
            # o hash da linha vai ser o mesmo e o Bloco 6 tem que deduplicar.
            #
            # E a unica sujeira que nao esta numa celula, e sim na linha a mais.
            # Por isso ela e contada aqui, e nao pela comparacao com 'antes'.
            linhas.append(dict(celulas))
            resumo.contar(sujeira)
        elif celulas != antes:
            resumo.contar(sujeira)
        else:
            resumo.contar_sem_efeito(sujeira)

    return linhas


# ---------------------------------------------------------------------------
# Escrita da planilha
# ---------------------------------------------------------------------------

def _escrever_titulo(
    aba: Worksheet,
    total_colunas: int,
    periodo: str,
    rng: random.Random,
    resumo: ResumoGeracao,
) -> int:
    """
    Escreve as linhas de titulo acima do cabecalho e devolve a linha do cabecalho.

    Sao duas ou tres, conforme a secao 4.2. Sem nome de empresa: o repositorio
    nao carrega nome inventado nem real.
    """
    titulos = [
        "RELATORIO DE APONTAMENTO DE PRODUCAO",
        f"Periodo: {periodo}",
    ]
    if rng.random() < 0.5:
        titulos.append("Emitido pelo sistema de apontamento - conferir antes de usar")

    for indice, texto in enumerate(titulos, start=1):
        celula = aba.cell(row=indice, column=1, value=texto)
        celula.font = Font(bold=(indice == 1), size=13 if indice == 1 else 10)
        aba.merge_cells(
            start_row=indice, start_column=1, end_row=indice, end_column=total_colunas
        )

    resumo.contar(SUJEIRA_CABECALHO)
    return len(titulos) + 1


def _escrever_rodape(
    aba: Worksheet,
    linha_atual: int,
    somas: dict[str, float],
    campos: list[str],
    rng: random.Random,
    resumo: ResumoGeracao,
) -> None:
    """Linhas em branco e a linha TOTAL, que nao e uma ordem de producao."""
    linha_total = linha_atual + rng.randint(1, 2)

    aba.cell(row=linha_total, column=1, value="TOTAL").font = Font(bold=True)
    for indice, campo in enumerate(campos, start=1):
        if campo in somas:
            celula = aba.cell(row=linha_total, column=indice, value=round(somas[campo], 1))
            celula.font = Font(bold=True)

    resumo.contar(SUJEIRA_RODAPE)


def _escrever_planilha(
    destino: Path,
    linhas: list[dict[str, Any]],
    somas: dict[str, float],
    cabecalhos: dict[str, str],
    periodo: str,
    rng: random.Random,
    resumo: ResumoGeracao,
) -> None:
    """Monta o .xlsx completo: titulo, cabecalho, linhas e rodape."""
    campos = list(cabecalhos)
    livro = Workbook()
    aba = livro.active
    aba.title = "Apontamento"

    linha_cabecalho = _escrever_titulo(aba, len(campos), periodo, rng, resumo)

    for indice, campo in enumerate(campos, start=1):
        celula = aba.cell(row=linha_cabecalho, column=indice, value=cabecalhos[campo])
        celula.font = Font(bold=True)
        celula.alignment = Alignment(horizontal="center", wrap_text=True)
        aba.column_dimensions[get_column_letter(indice)].width = max(
            12, min(34, len(cabecalhos[campo]) + 4)
        )

    linha_atual = linha_cabecalho
    for linha_atual, celulas in enumerate(linhas, start=linha_cabecalho + 1):
        for indice, campo in enumerate(campos, start=1):
            valor = celulas[campo]
            celula = aba.cell(row=linha_atual, column=indice, value=valor)
            if isinstance(valor, date):
                celula.number_format = "DD/MM/YYYY"

    _escrever_rodape(aba, linha_atual, somas, campos, rng, resumo)

    livro.save(destino)


def _dividir_em_arquivos(ordens: list[Ordem], arquivos: int) -> list[list[Ordem]]:
    """
    Divide as ordens em blocos cronologicos contiguos, um por arquivo.

    E como a fabrica exporta de verdade: um arquivo por periodo, nao um recorte
    aleatorio. Com 6 meses e 3 arquivos, cada arquivo cobre 2 meses.
    """
    if arquivos > len(ordens):
        raise ErroGeracao(
            f"pedidos {arquivos} arquivos, mas so ha {len(ordens)} ordens para "
            f"distribuir. Aumente 'gerador.ordens_por_mes' ou reduza --arquivos"
        )
    tamanho, resto = divmod(len(ordens), arquivos)
    blocos = []
    inicio = 0
    for indice in range(arquivos):
        # Os primeiros blocos ficam com uma ordem a mais quando a divisao nao e exata.
        fim = inicio + tamanho + (1 if indice < resto else 0)
        blocos.append(ordens[inicio:fim])
        inicio = fim
    return blocos


def _nome_do_arquivo(bloco: list[Ordem]) -> str:
    inicio = bloco[0].data_entrada
    fim = bloco[-1].data_entrada
    if (inicio.year, inicio.month) == (fim.year, fim.month):
        return f"producao_{inicio:%Y-%m}.xlsx"
    return f"producao_{inicio:%Y-%m}_a_{fim:%Y-%m}.xlsx"


def _escolher_cabecalhos(cfg: Configuracao, rng: random.Random) -> dict[str, str]:
    """
    Sorteia um apelido de coluna por campo, entre os aceitos no config.yaml.

    Cada arquivo sai com um conjunto de nomes de coluna, como se viesse de um
    setor diferente. E o que exercita de verdade o mapeamento por apelido do
    Bloco 3: se todo arquivo usasse sempre o primeiro nome, os outros apelidos
    do config nunca seriam testados.
    """
    return {campo: rng.choice(apelidos) for campo, apelidos in cfg.colunas.items()}


def _apurar_distribuicoes(ordens: list[Ordem], resumo: ResumoGeracao) -> None:
    """Apura status, OTD por mes e refugo por maquina sobre as ordens limpas."""
    for ordem in ordens:
        resumo.por_status[ordem.status] = resumo.por_status.get(ordem.status, 0) + 1

        if ordem.status == STATUS_ENTREGUE:
            mes = f"{ordem.data_entrada:%Y-%m}"
            no_prazo, total = resumo.otd_por_mes.get(mes, (0, 0))
            resumo.otd_por_mes[mes] = (
                no_prazo + (1 if ordem.data_entrega <= ordem.data_prazo else 0),
                total + 1,
            )

        boa, refugo = resumo.refugo_por_maquina.get(ordem.maquina, (0, 0))
        resumo.refugo_por_maquina[ordem.maquina] = (
            boa + ordem.tiragem_boa,
            refugo + ordem.refugo,
        )


# ---------------------------------------------------------------------------
# Orquestracao
# ---------------------------------------------------------------------------

def _conferir_status_conhecidos(cfg: Configuracao) -> None:
    """
    Confere que os tres status que o gerador sabe produzir existem no config.

    O gerador precisa saber o significado de cada status, nao so o nome, entao
    os tres sao constantes aqui. O que nao pode e o config renomear um deles e o
    gerador continuar escrevendo o nome antigo na planilha sem avisar.
    """
    conhecidos = (STATUS_EM_PRODUCAO, STATUS_ENTREGUE, STATUS_CANCELADA)
    faltando = [s for s in conhecidos if s not in cfg.status]
    if faltando:
        raise ErroGeracao(
            f"a secao 'status' do config nao declara {', '.join(repr(s) for s in faltando)}, "
            f"que o gerador precisa produzir. Declarados: "
            f"{', '.join(sorted(cfg.status))}"
        )


def gerar(
    cfg: Configuracao,
    meses: int | None = None,
    arquivos: int | None = None,
    data_fim: date = DATA_FIM_PADRAO,
    semente: int | None = None,
    destino: Path | None = None,
) -> ResumoGeracao:
    """
    Gera as planilhas sinteticas e devolve o resumo do que foi produzido.

    Os parametros sobrepoem o config.yaml quando informados, para atender
    'gerar-dados --meses 6 --arquivos 3' da secao 8.
    """
    meses = cfg.gerador["meses"] if meses is None else meses
    arquivos = cfg.gerador["arquivos"] if arquivos is None else arquivos
    semente = cfg.gerador["semente"] if semente is None else semente
    destino = cfg.caminho("entrada") if destino is None else Path(destino)

    if meses <= 0:
        raise ErroGeracao(f"--meses precisa ser maior que zero, recebido {meses}")
    if arquivos <= 0:
        raise ErroGeracao(f"--arquivos precisa ser maior que zero, recebido {arquivos}")

    _conferir_status_conhecidos(cfg)

    try:
        destino.mkdir(parents=True, exist_ok=True)
    except OSError as erro:
        raise ErroGeracao(
            f"nao foi possivel criar a pasta de destino {destino}: {erro}"
        ) from erro

    existentes = set(destino.glob("*.xlsx"))

    rng = random.Random(semente)
    resumo = ResumoGeracao()

    ordens, resumo.probabilidade_base_otd = gerar_ordens(cfg, meses, data_fim, rng)
    resumo.ordens_limpas = len(ordens)
    _apurar_distribuicoes(ordens, resumo)

    for bloco in _dividir_em_arquivos(ordens, arquivos):
        # As somas do rodape usam os valores limpos: e o total que o operador
        # teria visto antes de alguem estragar as celulas.
        somas = {campo: sum(getattr(o, campo) for o in bloco) for campo in CAMPOS_SOMADOS}
        linhas = sujar_linhas(bloco, cfg, rng, resumo)
        cabecalhos = _escolher_cabecalhos(cfg, rng)
        periodo = (
            f"{bloco[0].data_entrada:%d/%m/%Y} a {bloco[-1].data_entrada:%d/%m/%Y}"
        )
        caminho = destino / _nome_do_arquivo(bloco)

        try:
            _escrever_planilha(caminho, linhas, somas, cabecalhos, periodo, rng, resumo)
        except OSError as erro:
            raise ErroGeracao(f"nao foi possivel gravar {caminho}: {erro}") from erro

        resumo.arquivos.append(caminho)
        resumo.linhas_por_arquivo[caminho.name] = len(linhas)

    resumo.arquivos_preexistentes = sorted(existentes - set(resumo.arquivos))
    return resumo


def descrever_resumo(resumo: ResumoGeracao, cfg: Configuracao, destino: Path) -> str:
    """Monta o relatorio legivel da geracao."""
    linhas = [
        f"Destino: {destino}",
        f"Ordens limpas geradas: {resumo.ordens_limpas}",
        "",
        "Arquivos",
    ]
    for caminho in resumo.arquivos:
        linhas.append(
            f"  {caminho.name:<34} {resumo.linhas_por_arquivo[caminho.name]} linhas"
        )
    total_linhas = sum(resumo.linhas_por_arquivo.values())
    linhas.append(f"  {'TOTAL':<34} {total_linhas} linhas")

    linhas.append("")
    linhas.append("Sujeira injetada (conferida na celula depois de escrita)")
    for sujeira in SUJEIRAS_POR_LINHA + (SUJEIRA_CABECALHO, SUJEIRA_RODAPE):
        quantas = resumo.sujeiras.get(sujeira, 0)
        marca = " " if quantas else "!"
        sem_efeito = resumo.sujeiras_sem_efeito.get(sujeira, 0)
        nota = f"   ({sem_efeito} sorteada(s) sem efeito)" if sem_efeito else ""
        linhas.append(f" {marca} {sujeira:<34} {quantas}{nota}")

    total_sem_efeito = sum(resumo.sujeiras_sem_efeito.values())
    if total_sem_efeito:
        linhas.append("")
        linhas.append(
            f"  {total_sem_efeito} sorteio(s) de sujeira nao mudaram a celula e por isso "
            f"nao entram na contagem acima. E esperado: a variante sorteada pode ser o "
            f"proprio valor canonico."
        )

    ausentes = [s for s in SUJEIRAS_POR_LINHA if not resumo.sujeiras.get(s)]
    if ausentes:
        linhas.append("")
        linhas.append(
            f"ATENCAO: {len(ausentes)} tipo(s) de sujeira nao apareceram: "
            f"{', '.join(ausentes)}. Aumente 'gerador.taxa_sujeira' ou o volume."
        )

    linhas.append("")
    linhas.append("Status das ordens")
    for status in sorted(resumo.por_status):
        quantas = resumo.por_status[status]
        linhas.append(
            f"  {status:<16} {quantas:>5}  {quantas / resumo.ordens_limpas:>6.1%}"
        )

    concentracao = cfg.gerador["atrasos"]["concentracao"]
    linhas.append("")
    linhas.append(
        f"OTD por mes de entrada (meta {cfg.metas['otd_minimo_pct']}%, "
        f"alvo do gerador {cfg.gerador['atrasos']['otd_alvo_pct']}%)"
    )
    total_no_prazo = total_entregue = 0
    for mes in sorted(resumo.otd_por_mes):
        no_prazo, total = resumo.otd_por_mes[mes]
        total_no_prazo += no_prazo
        total_entregue += total
        marca = "  <- concentracao" if mes == concentracao["mes"] else ""
        linhas.append(
            f"  {mes}  {no_prazo:>4}/{total:<4} {no_prazo / total:>7.1%}{marca}"
        )
    linhas.append(
        f"  TOTAL    {total_no_prazo:>4}/{total_entregue:<4} "
        f"{total_no_prazo / total_entregue:>7.1%}"
    )
    linhas.append(
        f"  probabilidade base fora da concentracao: {resumo.probabilidade_base_otd:.1%}"
    )

    linhas.append("")
    linhas.append(f"Refugo por maquina (meta {cfg.metas['refugo_maximo_pct']}%)")
    for maquina in sorted(resumo.refugo_por_maquina):
        boa, refugo = resumo.refugo_por_maquina[maquina]
        pct = refugo / (boa + refugo) * 100
        marca = "acima da meta" if pct > cfg.metas["refugo_maximo_pct"] else "na meta"
        linhas.append(
            f"  {maquina:<24} {boa:>9} boas {refugo:>7} refugo  {pct:>5.2f}%  {marca}"
        )

    if resumo.arquivos_preexistentes:
        linhas.append("")
        linhas.append(
            "ATENCAO: a pasta ja tinha .xlsx que esta geracao nao sobrescreveu. "
            "A importacao vai le-los junto:"
        )
        for caminho in resumo.arquivos_preexistentes:
            linhas.append(f"  {caminho.name}")

    return "\n".join(linhas)


def _argumentos(argv: list[str]) -> argparse.Namespace:
    analisador = argparse.ArgumentParser(
        prog="python -m src.gerador_dados",
        description="Gera planilhas de producao sinteticas com sujeira proposital.",
    )
    analisador.add_argument("--config", default=str(CONFIG_PADRAO), help="caminho do config.yaml")
    analisador.add_argument("--meses", type=int, help="sobrepoe gerador.meses")
    analisador.add_argument("--arquivos", type=int, help="sobrepoe gerador.arquivos")
    analisador.add_argument("--semente", type=int, help="sobrepoe gerador.semente")
    analisador.add_argument(
        "--fim",
        default=DATA_FIM_PADRAO.isoformat(),
        help="ultimo dia da janela de historico, em AAAA-MM-DD",
    )
    analisador.add_argument("--destino", help="sobrepoe caminhos.entrada")
    return analisador.parse_args(argv)


def main(argv: list[str]) -> int:
    """Entrada de 'python -m src.gerador_dados'."""
    args = _argumentos(argv)

    try:
        data_fim = date.fromisoformat(args.fim)
    except ValueError as erro:
        print(
            f"ERRO: --fim {args.fim!r} nao e uma data AAAA-MM-DD ({erro})",
            file=sys.stderr,
        )
        return 1

    try:
        cfg = carregar(args.config)
    except ErroConfiguracao as erro:
        print(f"ERRO DE CONFIGURACAO: {erro}", file=sys.stderr)
        return CODIGO_SAIDA_ERRO_CONFIG

    destino = Path(args.destino) if args.destino else cfg.caminho("entrada")

    try:
        resumo = gerar(
            cfg,
            meses=args.meses,
            arquivos=args.arquivos,
            data_fim=data_fim,
            semente=args.semente,
            destino=destino,
        )
    except ErroGeracao as erro:
        print(f"ERRO DE GERACAO: {erro}", file=sys.stderr)
        return 1

    print(descrever_resumo(resumo, cfg, destino))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
